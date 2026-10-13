#!/usr/bin/env python3
"""
Generator script for Yonas_AI_DevOps_Philosophy_and_Tool_Interrogation.docx
Adheres strictly to the SE203 Week 7 Rubric naming convention:
'Your first name_AI: DevOps Philosophy and Tool Interrogation'
"""

import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

def build_docx(output_filename="Yonas_AI_DevOps_Philosophy_and_Tool_Interrogation.docx"):
    doc = docx.Document()

    # 1. Page Margins (1 inch on all sides)
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # 2. Document Title
    title = doc.add_paragraph()
    title_run = title.add_run('Yonas_AI: DevOps Philosophy and Tool Interrogation')
    title_run.font.name = 'Arial'
    title_run.font.size = Pt(20)
    title_run.font.bold = True
    title_run.font.color.rgb = RGBColor(15, 23, 42)
    title.paragraph_format.space_after = Pt(4)

    # 3. Subtitle / Academic Metadata
    meta = doc.add_paragraph()
    meta_run = meta.add_run(
        'Author: Yonas Leykun  |  Institution: Frontier Institute of Technology (SET)\n'
        'Course: SE 203: Application of Programming Concepts I  |  Module: Week 7 (DevOps Philosophy & Tool Interrogation)'
    )
    meta_run.font.name = 'Arial'
    meta_run.font.size = Pt(10)
    meta_run.font.italic = True
    meta_run.font.color.rgb = RGBColor(100, 116, 139)
    meta.paragraph_format.space_after = Pt(18)

    def add_heading(text):
        h = doc.add_paragraph()
        r = h.add_run(text)
        r.font.name = 'Arial'
        r.font.size = Pt(14)
        r.font.bold = True
        r.font.color.rgb = RGBColor(30, 58, 138)
        h.paragraph_format.space_before = Pt(14)
        h.paragraph_format.space_after = Pt(6)
        return h

    def add_body_p(text, bold_prefix=None):
        p = doc.add_paragraph()
        p.paragraph_format.line_spacing = 1.15
        p.paragraph_format.space_after = Pt(8)
        if bold_prefix:
            r_pre = p.add_run(bold_prefix)
            r_pre.font.name = 'Calibri'
            r_pre.font.size = Pt(11)
            r_pre.font.bold = True
            r_pre.font.color.rgb = RGBColor(15, 23, 42)
        r = p.add_run(text)
        r.font.name = 'Calibri'
        r.font.size = Pt(11)
        r.font.color.rgb = RGBColor(30, 41, 59)
        return p

    # Section 1: Prompt Formulation
    add_heading('1. Prompt Formulation')
    add_body_p(
        'To interrogate the core engineering philosophies underpinning Infrastructure as Code (IaC) and configuration management, '
        'the following prompt was formulated to engage the AI in the persona of a Senior DevOps Architect:'
    )

    prompt_text = (
        "My project asks me to write a Bash script and then a Puppet manifest to do the same job. This feels redundant.\n\n"
        "Act as a DevOps Architect and explain the 'why'. What is the fundamental difference in philosophy between an Imperative tool like Bash and a Declarative tool like Puppet?\n\n"
        "Why would a big company always choose Puppet over a Bash script for managing 1000 servers?"
    )

    prompt_p = doc.add_paragraph()
    prompt_p.paragraph_format.left_indent = Inches(0.4)
    prompt_p.paragraph_format.right_indent = Inches(0.4)
    prompt_p.paragraph_format.space_after = Pt(14)
    prompt_run = prompt_p.add_run(prompt_text)
    prompt_run.font.name = 'Consolas'
    prompt_run.font.size = Pt(9.5)
    prompt_run.font.color.rgb = RGBColor(15, 23, 42)

    # Section 2: Execution Screenshot
    add_heading('2. Execution Screenshot')
    add_body_p(
        'Below is the verified, uncropped screenshot of the AI assistant\'s complete architectural response. '
        'It captures the comprehensive comparative analysis between Imperative (Bash) and Declarative (Puppet) paradigms, '
        'the multi-server operational scalability breakdown, and the comparative summary matrix:'
    )

    img_path = os.path.join(os.path.dirname(__file__), 'screenshots', 'ai_response_screenshot.png')
    if not os.path.exists(img_path):
        img_path = 'screenshots/ai_response_screenshot.png'

    if os.path.exists(img_path):
        img_p = doc.add_paragraph()
        img_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        img_p.paragraph_format.space_after = Pt(4)
        img_p.add_run().add_picture(img_path, width=Inches(6.2))

        caption = doc.add_paragraph()
        caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap_run = caption.add_run(
            'Figure 1: Full AI assistant interaction showing the DevOps Architect synthesis of Imperative vs. Declarative philosophies, '
            'idempotence, multi-server operational scalability, and the architectural comparison matrix.'
        )
        cap_run.font.name = 'Arial'
        cap_run.font.size = Pt(9)
        cap_run.font.italic = True
        cap_run.font.color.rgb = RGBColor(100, 116, 139)
        caption.paragraph_format.space_after = Pt(14)
    else:
        add_body_p(f'[Image file not found at {img_path}]')

    # Section 3: Critical Reflection (Idempotence and Scalability)
    add_heading('3. Critical Reflection: Idempotence and Scalability')
    add_body_p(
        'In configuration management and distributed systems architecture, idempotence is the mathematical and operational '
        'property ensuring that an operation can be applied repeatedly without altering the state beyond the initial application. '
        'In contrast to imperative scripting where sequential commands execute blindly regardless of the machine\'s live state, '
        'an idempotent configuration engine such as Puppet evaluates the divergence between the live system state and the codified '
        'target state prior to initiating any mutation. When scaling across an enterprise fleet of 1,000 distributed nodes, servers '
        'inevitably fall victim to configuration drift—an entropy caused by uncoordinated hotfixes, manual administrator interventions, '
        'failed updates, or subtle environmental variances. Without idempotence, synchronizing a heterogeneous fleet via imperative '
        'scripts introduces catastrophic operational risks: reapplying a script risks appending redundant configuration directives, '
        're-triggering destructive initialization routines, or corrupting running daemons. Idempotent resource modeling transforms '
        'configuration application from an execution gamble into a safe, deterministic state reconciliation.',
        'Idempotence and Scalability: '
    )

    add_body_p(
        'Because idempotent declarations are inherently safe to execute indefinitely, they unlock the architectural foundation '
        'required for automated self-healing at enterprise scale. Rather than relying on manual, high-stress push deployments via SSH '
        'loops, an infrastructure managed by Puppet operates on a continuous, autonomous reconciliation loop where local agent daemons '
        'poll the centralized catalog and enforce compliance at regular intervals (e.g., every 30 minutes). When configuration drift '
        'occurs on any subset of the 1,000 servers—whether an engineer manually altered an Nginx virtual host file during an incident '
        'or an unauthorized daemon was launched—the idempotent agent detects the delta, rolls back unauthorized discrepancies, and '
        're-establishes the canonical baseline without human intervention. This continuous drift remediation guarantees fleet-wide '
        'consistency, prevents cascading configuration failures, and transforms configuration management from ephemeral command '
        'execution into an auditable, version-controlled single source of truth.',
        'Continuous Drift Remediation at Scale: '
    )

    # Section 4: GitHub Repository Update
    add_heading('4. GitHub Repository Update')
    add_body_p(
        'All task deliverables—including the core analysis text file (analysis.txt), imperative Bash script (0-install_nginx.sh), '
        'declarative Puppet manifest (0-install_nginx.pp), high-resolution execution screenshot, and comprehensive documentation—have '
        'been organized and committed to the public course repository:'
    )

    link_p1 = doc.add_paragraph()
    link_p1.paragraph_format.space_after = Pt(4)
    r1 = link_p1.add_run('• GitHub Repository: ')
    r1.bold = True
    r2 = link_p1.add_run('https://github.com/yonasleykun27/Artificial-Intelligence-in-Software-Engineering')
    r2.font.color.rgb = RGBColor(37, 99, 235)
    r2.underline = True

    link_p2 = doc.add_paragraph()
    link_p2.paragraph_format.space_after = Pt(14)
    r3 = link_p2.add_run('• Specific Task Folder: ')
    r3.bold = True
    r4 = link_p2.add_run('https://github.com/yonasleykun27/Artificial-Intelligence-in-Software-Engineering/tree/main/DevOps%20Philosophy%20and%20Tool%20Interrogation')
    r4.font.color.rgb = RGBColor(37, 99, 235)
    r4.underline = True

    target_path = os.path.join(os.path.dirname(__file__), output_filename)
    try:
        doc.save(target_path)
        print(f'Successfully generated {target_path}')
    except PermissionError:
        backup = target_path.replace('.docx', '_updated.docx')
        doc.save(backup)
        print(f'Original docx is locked. Saved to {backup} instead.')

if __name__ == '__main__':
    build_docx()
